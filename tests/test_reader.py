"""Tests for anp_fetcher.reader."""

import json
import zipfile
from pathlib import Path

import polars as pl

from anp_fetcher.reader import (
    decompress,
    normalize_dtypes,
    parse_filename,
    provenance_metadata,
    read_csv,
    read_excel,
    write_parquet,
)


def test_parse_filename_static():
    meta = parse_filename(Path("/data/anp/ie/ie-m3@20260529.xlsx"))
    assert meta["base_id"] == "ie-m3"
    assert meta["year"] is None
    assert meta["part"] is None
    assert meta["partition"] is None
    assert meta["modification"] == 20260529
    assert meta["ext"] == "xlsx"


def test_parse_filename_annual():
    meta = parse_filename(Path("vendas-municipais-asfalto_2020@20260101.xls"))
    assert meta["base_id"] == "vendas-municipais-asfalto"
    assert meta["year"] == 2020
    assert meta["part"] is None
    assert meta["partition"] == "2020"
    assert meta["modification"] == 20260101


def test_parse_filename_monthly():
    meta = parse_filename(Path("shpc-ca_2025-05@20250601.csv"))
    assert meta["base_id"] == "shpc-ca"
    assert meta["year"] == 2025
    assert meta["part"] == 5
    assert meta["partition"] == "2025-05"
    assert meta["ext"] == "csv"


def test_parse_filename_semestral():
    meta = parse_filename(Path("producao-terra-s1_2010-01@20260101.csv"))
    assert meta["base_id"] == "producao-terra-s1"
    assert meta["year"] == 2010
    assert meta["part"] == 1
    assert meta["partition"] == "2010-01"


def test_parse_filename_unrecognized():
    meta = parse_filename(Path("arquivo_estranho.csv"))
    assert meta["base_id"] == "arquivo_estranho"
    assert meta["partition"] is None
    assert meta["modification"] is None


def test_normalize_dtypes(tmp_path):
    df = pl.DataFrame(
        {
            "Valor de Venda": ["1.234,56", "2,5", "abc"],
            "Municipio": ["  SAO PAULO ", "  RIO  ", "  SP  "],
            "Ano": [2020, 2021, 2022],
        }
    )
    spec = {"numeric_columns": ["Valor de Venda"]}
    out = normalize_dtypes(df, spec)
    assert out["Valor de Venda"].dtype == pl.Float64
    assert out["Valor de Venda"].to_list() == [1234.56, 2.5, None]
    assert out["Municipio"].dtype == pl.Utf8
    assert out["Municipio"].to_list() == ["SAO PAULO", "RIO", "SP"]
    assert out["Ano"].dtype == pl.Int64


def test_read_csv_utf8(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("produto;valor\nGasolina;5,49\nEtanol;4,12\n", encoding="utf-8")
    df = read_csv(path, {"separator": ";", "encoding": "utf-8"})
    assert df.shape == (2, 2)
    # com infer_schema_length=0 o decimal vírgula permanece String no raw;
    # a conversão tipada é responsabilidade de normalize_dtypes (numeric_columns)
    assert df["valor"].dtype == pl.Utf8
    assert df["valor"][0] == "5,49"


def test_read_csv_latin1_fallback(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_bytes("produto;valor\nGasolina;5,49\n".encode("latin-1"))
    df = read_csv(path, {})  # encoding None -> fallback latin-1
    assert df.shape == (1, 2)
    assert df["produto"][0] == "Gasolina"


def test_read_excel_single_sheet(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append(["ano", "valor"])
    ws.append([2020, 10.5])
    path = tmp_path / "dados.xlsx"
    wb.save(path)
    df = read_excel(path, {})
    assert isinstance(df, pl.DataFrame)
    assert df.shape == (1, 2)


def test_read_excel_all_sheets(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Abas_2020"
    ws1.append(["ano", "valor"])
    ws1.append([2020, 10.5])
    ws2 = wb.create_sheet("Abas_2021")
    ws2.append(["ano", "valor"])
    ws2.append([2021, 20.5])
    path = tmp_path / "multi.xlsx"
    wb.save(path)
    frames = read_excel(path, {"sheet": "all"})
    assert isinstance(frames, dict)
    assert set(frames) == {"Abas_2020", "Abas_2021"}
    assert frames["Abas_2020"].shape == (1, 2)


def test_decompress_zip(tmp_path):
    raw = tmp_path / "producao-poco_2005@20260101.zip"
    inner = tmp_path / "producao_poco_2005.xls"
    inner.write_bytes(b"\xd0\xcf\x11\xe0 fake-xls")
    with zipfile.ZipFile(raw, "w") as zf:
        zf.write(inner, "producao_poco_2005.xls")
    extracted = decompress(raw)
    assert extracted.suffix == ".xls"
    assert extracted.exists()


def test_provenance_metadata(tmp_path):
    raw = tmp_path / "dados.csv"
    raw.write_text("a;b\n1;2\n")
    manifest = tmp_path / "dados.csv.manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "source_id": "anp",
                "dataset_id": "shpc-ca",
                "url": "https://example.org/x.csv",
                "sha256": "abc123",
                "fetched_at": "2026-01-01T00:00:00Z",
                "producer": "anp-fetcher",
            }
        )
    )
    meta = provenance_metadata(raw)
    assert meta["quantilica.dataset_id"] == "shpc-ca"
    assert meta["quantilica.origin_sha256"] == "abc123"


def test_provenance_metadata_missing(tmp_path):
    raw = tmp_path / "dados.csv"
    raw.write_text("a;b\n")
    assert provenance_metadata(raw) == {}


def test_write_parquet(tmp_path):
    df = pl.DataFrame({"a": [1, 2], "b": ["x", "y"]})
    target = tmp_path / "out" / "dados@20260101.parquet"
    result = write_parquet(df, target)
    assert result.exists()
    read = pl.read_parquet(target)
    assert read.shape == (2, 2)
